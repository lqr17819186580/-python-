#!/usr/bin/env python3
"""
run_full_pipeline_merge.py — merged version of run_full_pipeline.py and all dependencies.

Input files (must be in same dir):
  - 周业绩汇报PPT_20260404_更新.pptx  (or any template renamed to template.pptx)
  - S1-总览仪表盘.csv
  - S2-业务端视角.csv
  - S3-执行管理端.csv
  - S4-产品端视角.csv

Output:
  - 周业绩汇报PPT_FINAL.pptx

This script merges:
  - run_full_pipeline.py (main orchestrator)
  - data_loader.py (CSV parsing)
  - chart_xml_patch.py (XML-based chart updates)
  - chart_updates.py (chart data builders)
  - sowhat_slots.py (So What label discovery)
  - sowhat_generators.py (So What text generation)
  - apply_w14_patches.py (weekly patches)
  - final_gap_patch.py (final fixes)
  - update_ppt.py (main PPT update driver)
"""
import csv
import io
import re
import shutil
import sys
import zipfile
import os
from pathlib import Path
from copy import deepcopy
from dataclasses import dataclass, field
from typing import Optional

import pandas as pd
from lxml import etree

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

HERE = Path(__file__).parent

# ======== CONFIG ========
SRC_PPT_STAGE1 = "template.pptx"
OUT_PPT_STAGE1 = "周业绩汇报PPT_AUTO_UPDATED.pptx"
OUT_PPT_STAGE2 = "周业绩汇报PPT_AUTO_UPDATED_updated.pptx"
OUT_PPT_FINAL = "周业绩汇报PPT_FINAL.pptx"

CSV_S1 = "S1-总览仪表盘.csv"
CSV_S2 = "S2-业务端视角.csv"
CSV_S3 = "S3-执行管理端.csv"
CSV_S4 = "S4-产品端视角.csv"

TEMPLATE_SRC = HERE / "周业绩汇报PPT_20260404_更新.pptx"
if not TEMPLATE_SRC.exists():
    TEMPLATE_SRC = HERE / "template.pptx"

# ======== Monkey-patch: tolerate non-standard grouping='none' in line charts ========
import pptx.chart.plot as _plot_mod
from pptx.enum.chart import XL_CHART_TYPE as _XL
_orig_diff = _plot_mod.PlotTypeInspector._differentiate_line_chart_type
def _patched_line_type(cls, plot):
    try:
        return _orig_diff(plot)
    except (KeyError, Exception):
        has_none = bool(plot._element.xpath(
            'c:ser/c:marker/c:symbol[@val="none"]'))
        return _XL.LINE if has_none else _XL.LINE_MARKERS
_plot_mod.PlotTypeInspector._differentiate_line_chart_type = classmethod(_patched_line_type)

# ============================================================================
# Section 1: data_loader.py — Parse CSV files into DataFrames
# ============================================================================
HEADER_RE = re.compile(r"^([A-Z](?:-[^.\s]+)?)\.\s")

def _read_text_any_encoding(path: str) -> str:
    data = Path(path).read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "cp936"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")

def parse_blocks(path: str) -> dict:
    raw = _read_text_any_encoding(path).splitlines()
    blocks = {}
    cur_lines, cur_key = [], None
    for line in raw:
        stripped = line.lstrip("\ufeff")
        m = HEADER_RE.match(stripped)
        if m:
            if cur_key is not None:
                blocks[cur_key] = cur_lines
            cur_key = m.group(1)
            cur_lines = []
        elif cur_key is not None:
            cur_lines.append(line)
    if cur_key is not None:
        blocks[cur_key] = cur_lines

    out = {}
    for key, lines in blocks.items():
        data_lines = []
        for ln in lines:
            stripped = ln.strip()
            if not stripped:
                continue
            if stripped.startswith("📌"):
                continue
            first_char = stripped.lstrip('"').lstrip("'")[:1]
            if first_char in ("\ufffd", "?"):
                if stripped.count(",") < 4 or "\ufffd\ufffd" in stripped or "???" in stripped:
                    continue
            data_lines.append(ln)
        if len(data_lines) < 2:
            continue
        csv_text = "\n".join(data_lines)
        try:
            df = pd.read_csv(io.StringIO(csv_text), dtype=str, keep_default_na=False)
        except Exception:
            continue
        for c in df.columns:
            df[c] = df[c].astype(str).str.strip().str.strip('"')
        df.columns = [c.strip() for c in df.columns]
        df = df.loc[:, ~(df.columns.str.match(r"^Unnamed"))]
        out[key] = df
    return out

def num(v) -> float:
    if v is None:
        return 0.0
    s = str(v).strip().strip('"').replace(",", "").replace("%", "")
    if s in ("", "-", "—", "nan", "None"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0

def load_all(s1_path, s2_path, s3_path, s4_path):
    return (parse_blocks(s1_path), parse_blocks(s2_path),
            parse_blocks(s3_path), parse_blocks(s4_path))

# ============================================================================
# Section 2: chart_xml_patch.py — surgical update of chart values via XML
# ============================================================================
NS = {
    "c":  "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a":  "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r":  "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
C = NS["c"]

def _qn(tag):
    return f"{{{C}}}{tag}"

def _find_series(root):
    return root.findall(".//c:ser", NS)

def _rewrite_cache(cache_elem, new_values, numeric=True):
    for pt in list(cache_elem.findall("c:pt", NS)):
        cache_elem.remove(pt)
    pc = cache_elem.find("c:ptCount", NS)
    if pc is None:
        pc = etree.SubElement(cache_elem, _qn("ptCount"))
        cache_elem.insert(0, pc)
    pc.set("val", str(len(new_values)))
    for i, v in enumerate(new_values):
        pt = etree.SubElement(cache_elem, _qn("pt"))
        pt.set("idx", str(i))
        vel = etree.SubElement(pt, _qn("v"))
        if v is None:
            vel.text = ""
        elif numeric:
            vel.text = f"{float(v)}"
        else:
            vel.text = str(v)

def _rewrite_ref_with_cache(ref_elem, new_values, numeric=True):
    cache_name = "numCache" if numeric else "strCache"
    cache = ref_elem.find(f"c:{cache_name}", NS)
    if cache is None:
        cache = etree.SubElement(ref_elem, _qn(cache_name))
    _rewrite_cache(cache, new_values, numeric=numeric)

def _rewrite_series_values(ser, new_values):
    val = ser.find("c:val", NS)
    if val is None:
        return
    ref = val.find("c:numRef", NS)
    if ref is None:
        lit = val.find("c:numLit", NS)
        if lit is not None:
            _rewrite_cache(lit, new_values, numeric=True)
        return
    _rewrite_ref_with_cache(ref, new_values, numeric=True)

def _rewrite_series_categories(ser, new_cats):
    cat = ser.find("c:cat", NS)
    if cat is None:
        return
    ref = cat.find("c:strRef", NS)
    if ref is None:
        ref = cat.find("c:numRef", NS)
    if ref is None:
        lit = cat.find("c:strLit", NS)
        if lit is not None:
            _rewrite_cache(lit, new_cats, numeric=False)
        return
    numeric = (ref.tag == _qn("numRef"))
    _rewrite_ref_with_cache(ref, new_cats, numeric=numeric)

def _rewrite_series_name(ser, new_name):
    if new_name is None:
        return
    tx = ser.find("c:tx", NS)
    if tx is None:
        return
    ref = tx.find("c:strRef", NS)
    if ref is None:
        v = tx.find("c:v", NS)
        if v is not None:
            v.text = new_name
        return
    cache = ref.find("c:strCache", NS)
    if cache is None:
        cache = etree.SubElement(ref, _qn("strCache"))
    _rewrite_cache(cache, [new_name], numeric=False)

def patch_chart(chart_part, series_spec, categories=None, series_names=None):
    root = chart_part._element
    series = _find_series(root)
    n = min(len(series), len(series_spec))
    for i in range(n):
        ser = series[i]
        _rewrite_series_values(ser, series_spec[i])
        if categories is not None:
            _rewrite_series_categories(ser, categories)
        if series_names is not None and i < len(series_names):
            _rewrite_series_name(ser, series_names[i])

def patch_chart_dlbls(chart_part, series_index, labels):
    root = chart_part._element
    series = _find_series(root)
    if series_index >= len(series):
        return 0
    ser = series[series_index]
    dlbls = ser.findall("c:dLbls/c:dLbl", NS)
    hits = 0
    for dlbl in dlbls:
        idx_elem = dlbl.find("c:idx", NS)
        if idx_elem is None:
            continue
        try:
            idx = int(idx_elem.get("val"))
        except (TypeError, ValueError):
            continue
        if idx >= len(labels):
            continue
        new_text = labels[idx]
        tx = dlbl.find("c:tx", NS)
        if tx is None:
            continue
        rich = tx.find("c:rich", NS)
        if rich is None:
            continue
        p = rich.find("a:p", NS)
        if p is None:
            continue
        runs = p.findall("a:r", NS)
        if not runs:
            continue
        first_run = runs[0]
        first_t = first_run.find("a:t", NS)
        if first_t is None:
            first_t = etree.SubElement(first_run, f"{{{NS['a']}}}t")
        first_t.text = new_text
        for extra in runs[1:]:
            p.remove(extra)
        hits += 1
    return hits

# ============================================================================
# Section 3: sowhat_slots.py — discover So What labels
# ============================================================================
def find_sowhat_slots(prs):
    slots = []
    for si, s in enumerate(prs.slides):
        labels = [sh for sh in s.shapes
                  if sh.has_text_frame and sh.text_frame.text.strip().startswith('So What')]
        content_pool = []
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if len(t) < 40 or t.startswith('So What'):
                continue
            content_pool.append(sh)
        assigned = set()
        for sl in labels:
            rest = sl.text_frame.text.strip().replace('So What：', '', 1).strip()
            if len(rest) > 20:
                slots.append({'slide': si, 'label': sl, 'content': sl, 'type': 'MERGED'})
                continue
            lcx = sl.left + (sl.width or 0) // 2
            ly = sl.top
            best = None
            best_score = None
            for sh in content_pool:
                if id(sh) in assigned:
                    continue
                cy = sh.top
                if cy < ly - 20000 or cy > ly + 400000:
                    continue
                cx = sh.left + (sh.width or 0) // 2
                left_edge = sh.left
                right_edge = sh.left + (sh.width or 0)
                contains = left_edge <= lcx <= right_edge
                if not contains and abs(cx - lcx) > 3_000_000:
                    continue
                score = (cy - ly) * 5
                if not contains:
                    score += abs(cx - lcx)
                if best_score is None or score < best_score:
                    best = sh
                    best_score = score
            if best is not None:
                slots.append({'slide': si, 'label': sl, 'content': best, 'type': 'SEPARATE'})
                assigned.add(id(best))
            else:
                slots.append({'slide': si, 'label': sl, 'content': None, 'type': 'ORPHAN'})
    return slots

# ============================================================================
# Section 4: chart_updates.py — build CategoryChartData for charts
# ============================================================================
def _chart_row(df, col_match_value, col_name):
    first_col = df.columns[0]
    sub = df[df[first_col] == col_match_value]
    if len(sub) == 0:
        return 0.0
    return num(sub.iloc[0][col_name])

def _chart_total_row(df):
    first_col = df.columns[0]
    sub = df[df[first_col] == "合计"]
    return sub.iloc[0] if len(sub) else None

def chart_to_m(v):
    return round(num(v) / 1_000_000, 2)

def chart_to_w(v):
    return round(num(v) / 10_000, 2)

def slide1_chart_business_type(S1):
    g = S1["G"]
    cats = ["经代业务", "代理人业务", "KA 业务"]
    mapping = {"经代业务": "经代业务", "代理人业务": "代理人业务", "KA 业务": "KA业务"}
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("批核 APE",
                  [chart_to_m(_chart_row(g, mapping[c], "2026批核APE")) for c in cats])
    cd.add_series("未批核 APE",
                  [chart_to_m(_chart_row(g, mapping[c], "未批核APE")) for c in cats])
    cd.add_series("待签 APE",
                  [chart_to_m(_chart_row(g, mapping[c], "待签APE")) for c in cats])
    return cd

def slide1_chart_monthly_trend(S1):
    cC, cD, cE = S1["C"], S1["D"], S1["E"]
    months = [m for m in cC.iloc[:, 0] if m and m.startswith("202")]
    def label(ym):
        y, m = ym.split("-")
        mn = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
        return f"{y[-2:]}-{mn[int(m)-1]}"
    cats = [label(m) for m in months]
    def col_for(df, ym, col):
        row = df[df.iloc[:, 0] == ym]
        return chart_to_m(row.iloc[0][col]) if len(row) else 0.0
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("预约 APE(M)", [col_for(cC, m, "APE") for m in months])
    cd.add_series("签单 APE(M)", [col_for(cD, m, "APE") for m in months])
    cd.add_series("批核 APE(M)", [col_for(cE, m, "APE") for m in months])
    return cd

def _slide2_bars(S1, block_letter, series_name):
    df = S1[block_letter]
    df26 = df[df.iloc[:, 0].astype(str).str.startswith("2026")]
    cats = ["1月", "2月", "3月", "4月"]
    vals = [0.0] * 4
    for _, r in df26.iterrows():
        m_idx = int(r.iloc[0].split("-")[1]) - 1
        if 0 <= m_idx < 4:
            vals[m_idx] = chart_to_m(r["APE"])
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(series_name, vals)
    return cd, vals

def slide2_appointment(S1):
    cd, _ = _slide2_bars(S1, "C", "预约 APE(M)")
    return cd

def slide2_signed(S1):
    cd, _ = _slide2_bars(S1, "D", "签单 APE(M)")
    return cd

def slide2_approved(S1):
    cd, _ = _slide2_bars(S1, "E", "批核 APE(M)")
    return cd

def slide2_forecast_chart(S1):
    cE = S1["E"]
    actuals = [0.0] * 12
    for _, r in cE.iterrows():
        ym = str(r.iloc[0])
        if ym.startswith("2026-"):
            m_idx = int(ym.split("-")[1]) - 1
            if 0 <= m_idx < 12:
                actuals[m_idx] = chart_to_m(r["APE"])
    forecast = [None, None, None, None, 80.0, 88.0, 95.0, 100.0, 100.0, 100.0, 100.0, 96.0]
    cd = CategoryChartData()
    cd.categories = ["1月","2月","3月","4月","5月","6月","7月","8月","9月","10月","11月","12月"]
    cd.add_series("实际批核 APE(M)",
                  [v if i < 4 else None for i, v in enumerate(actuals)])
    cd.add_series("预测 APE(M)", forecast)
    return cd

def slide3_sunlife_trend(S2):
    fA, gA, hA = S2["F-APE"], S2["G-APE"], S2["H-APE"]
    cats = ["26-Jan", "26-Feb", "26-Mar", "26-Apr"]
    months = ["2026-01", "2026-02", "2026-03", "2026-04"]
    def total(df, col):
        row = df[df.iloc[:, 0] == "合计"]
        return chart_to_m(row.iloc[0][col]) if len(row) else 0.0
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("预约 APE(M)", [total(fA, m) for m in months])
    cd.add_series("签单 APE(M)", [total(gA, m) for m in months])
    cd.add_series("批核 APE(M)", [total(hA, m) for m in months])
    return cd

def slide4_channel_trend(S2, channel_name):
    cA, cD, cE = S2["C-APE"], S2["D-APE"], S2["E-APE"]
    months = ["2026-01", "2026-02", "2026-03", "2026-04"]
    def val(df, col):
        row = df[df.iloc[:, 0] == channel_name]
        return chart_to_m(row.iloc[0][col]) if len(row) else 0.0
    cd = CategoryChartData()
    cd.categories = ["Jan", "Feb", "Mar", "Apr"]
    cd.add_series("预约", [val(cA, m) for m in months])
    cd.add_series("签单", [val(cD, m) for m in months])
    cd.add_series("批核", [val(cE, m) for m in months])
    return cd

def slide5_donut(S2):
    total = _chart_total_row(S2["A"])
    cd = CategoryChartData()
    cd.categories = ["已批核", "未批核", "待签"]
    cd.add_series("APE构成",
                  [chart_to_m(total["2026批核APE"]),
                   chart_to_m(total["未批核APE"]),
                   chart_to_m(total["待签APE"])])
    return cd

def slide5_target_vs_actual(S2):
    df = S2["A"]
    df = df[df["业务细分"].isin([
        "永明经代", "天领业务", "BK业务", "合伙转介业务",
        "成事家办", "同行经代", "ICLUB业务", "IFA业务"
    ])]
    name_map = {
        "永明经代": "永明经代", "天领业务": "天领业务", "BK业务": "BK业务",
        "合伙转介业务": "合伙转介", "成事家办": "成事家办", "同行经代": "同行经代",
        "ICLUB业务": "ICLUB", "IFA业务": "IFA业务",
    }
    order = ["永明经代", "天领业务", "BK业务", "合伙转介业务",
             "成事家办", "同行经代", "ICLUB业务", "IFA业务"]
    cats = [name_map[k] for k in order]
    issued, unbat, pend, gap = [], [], [], []
    for k in order:
        row = df[df["业务细分"] == k].iloc[0]
        i = chart_to_m(row["2026批核APE"])
        u = chart_to_m(row["未批核APE"])
        p = chart_to_m(row["待签APE"])
        t = chart_to_m(row["目标APE"])
        issued.append(i); unbat.append(u); pend.append(p)
        gap.append(round(t - (i + u + p), 2))
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("已批核 APE(M)", issued)
    cd.add_series("未批核 APE(M)", unbat)
    cd.add_series("待签 APE(M)",   pend)
    cd.add_series("目标缺口 APE(M)", gap)
    return cd

def slide6_weekly_trend(S3):
    a = S3["A-APE"]
    week_cols = [c for c in a.columns if c.startswith("2026W")]
    cats = [c.replace("2026", "") for c in week_cols]
    cd = CategoryChartData()
    cd.categories = cats
    for stage in ["预约", "签单", "递交", "批核"]:
        row = a[a.iloc[:, 0] == stage].iloc[0]
        cd.add_series(stage, [chart_to_m(row[c]) for c in week_cols])
    return cd

def slide8_referrer(S2):
    j = S2["J"]
    j = j[j.iloc[:, 0] != "合计"]
    j = j.copy()
    j["_sort"] = j["2026批核APE"].apply(num)
    j = j.sort_values("_sort", ascending=False)
    cats = j.iloc[:, 0].tolist()
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("批核APE(万)",   [chart_to_w(v) for v in j["2026批核APE"]])
    cd.add_series("未批核APE(万)", [chart_to_w(v) for v in j["未批核APE"]])
    cd.add_series("待签APE(万)",   [chart_to_w(v) for v in j["待签APE"]])
    return cd

def slide8_top10_ka(S2):
    k = S2["K"]
    k = k[k.iloc[:, 0] != "合计"].copy()
    k["_sort"] = k["2026批核APE"].apply(num)
    k = k.sort_values("_sort", ascending=False).head(10)
    cats = k.iloc[:, 0].tolist()
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("批核APE(万)",   [chart_to_w(v) for v in k["2026批核APE"]])
    cd.add_series("未批核APE(万)", [chart_to_w(v) for v in k["未批核APE"]])
    cd.add_series("待签APE(万)",   [chart_to_w(v) for v in k["待签APE"]])
    return cd

def slide9_peer_weekly(S3):
    j, k, l = S3["J-APE"], S3["K-APE"], S3["L-APE"]
    week_cols = [c for c in j.columns if c.startswith("2026W")]
    cats = [c.replace("2026", "") for c in week_cols]
    def totals(df):
        row = df[df.iloc[:, 0] == "合计"].iloc[0]
        return [chart_to_m(row[c]) for c in week_cols]
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("同行预约(M)", totals(j))
    cd.add_series("同行签单(M)", totals(k))
    cd.add_series("同行批核(M)", totals(l))
    return cd

def _slide10_bank_monthly(df):
    _MO_ZH = {1:"1月",2:"2月",3:"3月",4:"4月",5:"5月",6:"6月",
              7:"7月",8:"8月",9:"9月",10:"10月",11:"11月",12:"12月"}
    months = sorted([c for c in df.columns if c.startswith("2026-")])
    cats = [_MO_ZH.get(int(c[5:7]), c) for c in months]
    def val(bank, month):
        row = df[df.iloc[:, 0] == bank]
        return chart_to_m(row.iloc[0][month]) if len(row) else 0.0
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("民生银行", [val("民生银行", m) for m in months])
    cd.add_series("平安银行", [val("平安银行", m) for m in months])
    return cd

def slide10_bk_appointment(S2):   return _slide10_bank_monthly(S2["P-APE"])
def slide10_bk_signed(S2):        return _slide10_bank_monthly(S2["Q-APE"])
def slide10_bk_approved(S2):      return _slide10_bank_monthly(S2["R-APE"])

def slide10_bk_ka(S2):
    o = S2["O"]
    cats = ["民生银行", "平安银行"]
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("批核APE(M)",   [chart_to_m(_chart_row(o, c, "2026批核APE")) for c in cats])
    cd.add_series("未批核APE(M)", [chart_to_m(_chart_row(o, c, "未批核APE"))   for c in cats])
    cd.add_series("待签APE(M)",   [chart_to_m(_chart_row(o, c, "待签APE"))     for c in cats])
    return cd

def slide10_donut_target(S2):
    bk = S2["A"][S2["A"]["业务细分"] == "BK业务"].iloc[0]
    issued = chart_to_m(bk["2026批核APE"])
    remain = chart_to_m(bk["目标APE"]) - issued
    cd = CategoryChartData()
    cd.categories = ["已批核", "目标剩余"]
    cd.add_series("APE构成", [issued, remain])
    return cd

def slide11_bank_weekly(S3):
    m, n, o = S3["M-APE"], S3["N-APE"], S3["O-APE"]
    week_cols = [c for c in m.columns if c.startswith("2026W")]
    cats = [c.replace("2026", "") for c in week_cols]
    def totals(df):
        row = df[df.iloc[:, 0] == "合计"].iloc[0]
        return [chart_to_m(row[c]) for c in week_cols]
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("银行预约(M)", totals(m))
    cd.add_series("银行签单(M)", totals(n))
    cd.add_series("银行批核(M)", totals(o))
    return cd

def slide11_branch_ranking(S2):
    s = S2["S-APE"]
    s = s[s.iloc[:, 0] != "合计"].copy()
    s["_sort"] = s["合计"].apply(num)
    s = s.sort_values("_sort", ascending=False)
    s = s[s["_sort"] > 0]
    cats = s.iloc[:, 0].tolist()
    vals = [chart_to_m(v) for v in s["合计"]]
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("批核APE(M)", vals)
    return cd