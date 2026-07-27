from pptx import Presentation

target_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL.pptx"
template_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\template.pptx"

print("=== Target PPT Text 29 ===")
prs = Presentation(target_path)
slide = prs.slides[4]
for s in slide.shapes:
    if s.name == "Text 29":
        print(f"  name={s.name}, top={s.top}, left={s.left}, width={s.width}, height={s.height}")
        print(f"  text='{s.text_frame.text.strip()}'")

print("\n=== Template PPT Text 29 ===")
template_prs = Presentation(template_path)
template_slide = template_prs.slides[4]
for s in template_slide.shapes:
    if s.name == "Text 29":
        print(f"  name={s.name}, top={s.top}, left={s.left}, width={s.width}, height={s.height}")
        print(f"  text='{s.text_frame.text.strip()}'")

print("\n=== Target PPT Text 129 ===")
for s in slide.shapes:
    if s.name == "Text 129":
        print(f"  name={s.name}, top={s.top}, left={s.left}, width={s.width}, height={s.height}")
        print(f"  text='{s.text_frame.text.strip()}'")

print("\n=== Template PPT Text 129 ===")
for s in template_slide.shapes:
    if s.name == "Text 129":
        print(f"  name={s.name}, top={s.top}, left={s.left}, width={s.width}, height={s.height}")
        print(f"  text='{s.text_frame.text.strip()}'")
