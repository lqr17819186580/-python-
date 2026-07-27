from pptx import Presentation

def extract_slide_info(prs, slide_idx):
    if slide_idx < 1 or slide_idx > len(prs.slides):
        return None
    slide = prs.slides[slide_idx-1]
    info = {
        'index': slide_idx,
        'title': '',
        'texts': [],
        'charts': [],
        'tables': []
    }
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                info['texts'].append(text)
                if not info['title']:
                    info['title'] = text[:50]
        if shape.has_chart:
            chart = shape.chart
            chart_info = {
                'name': shape.name,
                'categories': [],
                'series': []
            }
            try:
                if hasattr(chart, 'categories'):
                    chart_info['categories'] = [str(c.label) for c in chart.categories]
                for series in chart.series:
                    chart_info['series'].append({
                        'name': series.name if series.name else '未命名',
                        'values': [str(v) for v in series.values]
                    })
            except:
                pass
            info['charts'].append(chart_info)
        if shape.has_table:
            table = shape.table
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip()[:50] for cell in row.cells]
                table_data.append(row_data)
            info['tables'].append(table_data)
    return info

ref_ppt = Presentation('周业绩汇报PPT_W28_20260717.pptx')
test_ppt = Presentation('周业绩汇报PPT_FINAL.pptx')

pages = [1, 4, 5, 7, 8, 9]
for page in pages:
    print('=== Page {} ==='.format(page))
    ref_info = extract_slide_info(ref_ppt, page)
    test_info = extract_slide_info(test_ppt, page)
    
    if ref_info:
        print('参考标题: {}'.format(ref_info['title']))
        print('参考图表数: {}'.format(len(ref_info['charts'])))
        for i, ch in enumerate(ref_info['charts']):
            print('  图表{}: {}'.format(i, ch['name']))
            print('    分类: {}'.format(ch['categories']))
            for s in ch['series']:
                print('    系列: {}'.format(s['name']))
        print('参考表格数: {}'.format(len(ref_info['tables'])))
    else:
        print('参考文件无此页')
    
    if test_info:
        print('测试标题: {}'.format(test_info['title']))
        print('测试图表数: {}'.format(len(test_info['charts'])))
        for i, ch in enumerate(test_info['charts']):
            print('  图表{}: {}'.format(i, ch['name']))
            print('    分类: {}'.format(ch['categories']))
            for s in ch['series']:
                print('    系列: {}'.format(s['name']))
        print('测试表格数: {}'.format(len(test_info['tables'])))
    else:
        print('测试文件无此页')
    
    print()
