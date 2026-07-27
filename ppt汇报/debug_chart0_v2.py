import pptx

prs = pptx.Presentation('周业绩汇报PPT_FINAL.pptx')
slide = prs.slides[0]

for shape in slide.shapes:
    if shape.name == 'Chart 0':
        chart_part = shape.chart.part
        root = chart_part._element
        
        print('=== Chart 0 XML Structure (Template) ===')
        print()
        
        barCharts = root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}barChart')
        print(f'barChart elements: {len(barCharts)}')
        
        for j, bc in enumerate(barCharts):
            print(f'\n--- barChart {j} ---')
            
            grouping = bc.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}grouping')
            print(f'Grouping: {grouping.get("val") if grouping is not None else "None"}')
            
            series_in_bc = bc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}ser')
            print(f'Series in barChart {j}: {len(series_in_bc)}')
            
            for i, ser in enumerate(series_in_bc):
                tx = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
                name = tx.text if tx is not None else 'Unknown'
                
                val = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}val')
                if val is not None:
                    numRef = val.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numRef')
                    numLit = val.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numLit')
                    
                    if numRef is not None:
                        f = numRef.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}f')
                        f_text = f.text if f is not None else 'None'
                        nc = numRef.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache')
                        pts = nc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt') if nc is not None else []
                        print(f'  Series {i}: {name}')
                        print(f'    numRef/f: {f_text}')
                        print(f'    numCache points: {len(pts)}')
                    
                    if numLit is not None:
                        nc = numLit.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache')
                        pts = nc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt') if nc is not None else []
                        print(f'  Series {i}: {name} (numLit)')
                        print(f'    numCache points: {len(pts)}')

        print('\n=== All ser elements in document ===')
        all_series = root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}ser')
        print(f'Total ser elements: {len(all_series)}')
        for i, ser in enumerate(all_series):
            tx = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
            name = tx.text if tx is not None else 'Unknown'
            print(f'  Ser {i}: {name}')
