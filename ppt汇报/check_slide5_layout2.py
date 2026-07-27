from pptx import Presentation
from pptx.util import Emu

prs = Presentation('周业绩汇报PPT_FINAL.pptx')
slide = prs.slides[4]

print('Slide 5 shapes near J and K sections:')
for s in slide.shapes:
    if s.has_text_frame:
        text = s.text_frame.text.strip()[:30]
        if '保单阶段' in text or '保单总量' in text:
            print(f'  {s.name}: "{text}" - top={s.top}, bottom={s.top + s.height}')

print('\nSlide 5 all text shapes with top position (sorted):')
all_text = []
for s in slide.shapes:
    if s.has_text_frame and s.top is not None:
        text = s.text_frame.text.strip()[:20]
        all_text.append((s.top, s.name, text))

all_text.sort()
for top, name, text in all_text:
    print(f'  top={top}: {name} - "{text}"')