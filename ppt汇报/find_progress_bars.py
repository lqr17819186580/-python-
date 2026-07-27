from pptx import Presentation

template_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\template.pptx"
target_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL.pptx"

print("=== Analyzing template.pptx K section shapes ===")
template_prs = Presentation(template_path)
template_slide = template_prs.slides[4]

for s in template_slide.shapes:
    if s.top is not None and 850000 <= s.top <= 2800000:
        fill_color = None
        line_color = None
        has_text = s.has_text_frame
        text = s.text_frame.text.strip()[:40] if has_text else ""
        
        if hasattr(s, 'fill'):
            try:
                fill = s.fill
                if fill.type == 1:
                    fill_color = fill.fore_color.rgb
                elif fill.type == 0:
                    fill_color = "NoFill"
            except:
                pass
        
        print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, "
              f"has_text={has_text}, fill={fill_color}, text='{text}'")

print("\n\n=== Analyzing target PPT K section shapes ===")
prs = Presentation(target_path)
slide = prs.slides[4]

for s in slide.shapes:
    if s.top is not None and 850000 <= s.top <= 2800000:
        fill_color = None
        line_color = None
        has_text = s.has_text_frame
        text = s.text_frame.text.strip()[:40] if has_text else ""
        
        if hasattr(s, 'fill'):
            try:
                fill = s.fill
                if fill.type == 1:
                    fill_color = fill.fore_color.rgb
                elif fill.type == 0:
                    fill_color = "NoFill"
            except:
                pass
        
        print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, "
              f"has_text={has_text}, fill={fill_color}, text='{text}'")
