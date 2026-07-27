import pptx
from chart_xml_patch import _find_series

prs = pptx.Presentation('周业绩汇报PPT_W28_20260717.pptx')
slide = prs.slides[0]
for shape in slide.shapes:
    if shape.name == 'Chart 0' and shape.has_chart:
        chart_part = shape.chart.part
        series = _find_series(chart_part._element)
        print('参考PPT Chart 0:')
        for i, ser in enumerate(series):
            tx = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
            name = tx.text if tx is not None else 'Unknown'
            nc = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache')
            pts = nc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt') if nc is not None else []
            print(f'  Series {i}: \"{name}\", points: {len(pts)}')
        break