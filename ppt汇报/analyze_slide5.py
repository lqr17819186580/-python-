from pptx import Presentation
from pptx.util import Emu

target_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL.pptx"
template_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\template.pptx"

print("=== Analyzing target PPT Slide 5 shapes ===")
prs = Presentation(target_path)
slide = prs.slides[4]

print("\n--- All shapes with position info ---")
for s in slide.shapes:
    if s.top is not None:
        text_content = ""
        if s.has_text_frame:
            text_content = s.text_frame.text.strip()[:50]
        print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, has_text={s.has_text_frame}, text='{text_content}'")

print("\n--- K section shapes (So What area) ---")
for s in slide.shapes:
    if s.top is not None and 2800000 <= s.top <= 3500000:
        text_content = ""
        if s.has_text_frame:
            text_content = s.text_frame.text.strip()[:50]
        print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, text='{text_content}'")

print("\n--- L section title shapes ---")
for s in slide.shapes:
    if s.top is not None and 3300000 <= s.top <= 4000000:
        text_content = ""
        if s.has_text_frame:
            text_content = s.text_frame.text.strip()[:50]
        print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, text='{text_content}'")

print("\n=== Comparing with template PPT ===")
template_prs = Presentation(template_path)
template_slide = template_prs.slides[4]

print("\n--- Template K section shapes (So What area) ---")
for s in template_slide.shapes:
    if s.top is not None and 2800000 <= s.top <= 3500000:
        text_content = ""
        if s.has_text_frame:
            text_content = s.text_frame.text.strip()[:50]
        print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, text='{text_content}'")

print("\n--- Template L section title shapes ---")
for s in template_slide.shapes:
    if s.top is not None and 3300000 <= s.top <= 4000000:
        text_content = ""
        if s.has_text_frame:
            text_content = s.text_frame.text.strip()[:50]
        print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, text='{text_content}'")

print("\n=== Progress bar shapes comparison ===")
print("\n--- Target progress bars (around K section) ---")
for s in slide.shapes:
    if s.top is not None and 850000 <= s.top <= 2800000:
        if s.name.startswith('Shape') and not s.has_text_frame:
            fill_color = None
            if hasattr(s, 'fill'):
                try:
                    fill = s.fill
                    if fill.type == 1:
                        fill_color = fill.fore_color.rgb
                except:
                    pass
            print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, fill={fill_color}")

print("\n--- Template progress bars (around K section) ---")
for s in template_slide.shapes:
    if s.top is not None and 850000 <= s.top <= 2800000:
        if s.name.startswith('Shape') and not s.has_text_frame:
            fill_color = None
            if hasattr(s, 'fill'):
                try:
                    fill = s.fill
                    if fill.type == 1:
                        fill_color = fill.fore_color.rgb
                except:
                    pass
            print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, fill={fill_color}")
