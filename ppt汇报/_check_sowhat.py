from pptx import Presentation
from pptx.util import Emu

# 检查template.pptx第7页的所有形状
prs = Presentation('template.pptx')
slide7 = prs.slides[6]
print('=== template.pptx 第7页 所有形状 ===')
print(f'总形状数: {len(slide7.shapes)}')
print()

so_what_shapes = []
for i, shape in enumerate(slide7.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        # 找出So What相关文本框（通常包含较长的分析文案）
        if len(text) > 100 and shape.top is not None:
            so_what_shapes.append((i, shape.name, text, shape.top, shape.left, shape.width, shape.height))
            print(f'形状{i}: {shape.name}')
            print(f'  位置: top={shape.top}, left={shape.left}, 宽={shape.width}, 高={shape.height}')
            print(f'  文本({len(text)}字): {text[:100]}...')
            print()

print(f'\n共找到 {len(so_what_shapes)} 个长文本框（So What）')

# 检查FINAL.pptx
print()
print('=== FINAL.pptx 第7页 So What对比 ===')
prs2 = Presentation('周业绩汇报PPT_FINAL.pptx')
slide7_final = prs2.slides[6]
final_sw = []
for i, shape in enumerate(slide7_final.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if len(text) > 100 and shape.top is not None:
            final_sw.append((i, shape.name, text, shape.top, shape.left, shape.width, shape.height))

print(f'FINAL中找到 {len(final_sw)} 个长文本框')
print()

# 对比
template_texts = set(t[2][:80] for t in so_what_shapes)
final_texts = set(t[2][:80] for t in final_sw)

missing = template_texts - final_texts
if missing:
    print('缺失的So What:')
    for m in missing:
        print(f'  - {m}...')
else:
    print('没有缺失的So What')

# 显示FINAL中的So What
print('\nFINAL中的So What:')
for i, name, text, top, left, w, h in final_sw:
    print(f'  形状{i}: top={top}, left={left}')
    print(f'    {text[:80]}...')