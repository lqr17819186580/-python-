from pptx import Presentation

def analyze_slide8(ppt_path, label):
    print(f"\n{'='*60}")
    print(f"Slide 8 Analysis: {label}")
    print(f"Path: {ppt_path}")
    print(f"{'='*60}")
    
    prs = Presentation(ppt_path)
    if len(prs.slides) < 8:
        print(f"  ERROR: Only {len(prs.slides)} slides found")
        return
    
    slide = prs.slides[7]  # 0-indexed, slide 8 is index 7
    print(f"\n  Total shapes: {len(slide.shapes)}")
    
    for i, shape in enumerate(slide.shapes):
        name = shape.name
        shape_type = shape.shape_type
        has_text = shape.has_text_frame
        has_table = shape.has_table
        has_chart = shape.has_chart
        left = shape.left if hasattr(shape, 'left') else None
        top = shape.top if hasattr(shape, 'top') else None
        width = shape.width if hasattr(shape, 'width') else None
        height = shape.height if hasattr(shape, 'height') else None
        
        text_content = ""
        if has_text:
            text_content = shape.text[:50] if shape.text else ""
        
        print(f"\n  Shape [{i}]: name='{name}' type={shape_type}")
        print(f"    has_text={has_text}, has_table={has_table}, has_chart={has_chart}")
        if left is not None:
            print(f"    position: left={left}, top={top}, width={width}, height={height}")
        if text_content:
            print(f"    text: {text_content}...")
        if has_chart:
            chart = shape.chart
            print(f"    chart type: {chart.chart_type}")
            try:
                cat = chart.categories
                print(f"    categories: {list(cat)}")
                for series in chart.series:
                    print(f"    series: {series.name}, values: {len(series.values)}")
            except Exception as e:
                print(f"    chart data error: {e}")
        if has_table:
            table = shape.table
            print(f"    table: {len(table.rows)} rows x {len(table.columns)} cols")
            # Print header
            headers = []
            for col in range(len(table.columns)):
                headers.append(table.cell(0, col).text)
            print(f"    headers: {headers}")

if __name__ == "__main__":
    analyze_slide8("template.pptx", "TEMPLATE")
    analyze_slide8("周业绩汇报PPT_FINAL.pptx", "FINAL")