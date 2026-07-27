from pptx import Presentation
from pptx.util import Emu

target_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL.pptx"
output_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL_fixed.pptx"

print("=== Fixing Slide 5 layout issues ===\n")

prs = Presentation(target_path)
slide = prs.slides[4]

print("--- 1. Fixing K section So-What shapes ---")
so_what_positions = {
    "Shape 126": {"top": 2902077, "left": 4608576, "width": 7461504, "height": 384048},
    "Shape 127": {"top": 2902077, "left": 4608576, "width": 7461504, "height": 27432},
    "Text 128": {"top": 2938653, "left": 4700016, "width": 7296912, "height": 128016},
    "Text 129": {"top": 3057525, "left": 4700016, "width": 7296912, "height": 210312},
}

for s in slide.shapes:
    if s.name in so_what_positions:
        pos = so_what_positions[s.name]
        print(f"  {s.name}: top={s.top} -> {pos['top']}")
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']

print("\n--- 2. Fixing K section progress bars ---")
progress_bar_positions = {
    "Shape 98": {"top": 1821942, "left": 7333488, "width": 1280160, "height": 146304},
    "Shape 99": {"top": 1821942, "left": 7333488, "width": 839785, "height": 146304},
    "Shape 109": {"top": 1821942, "left": 8906256, "width": 1280160, "height": 146304},
    "Shape 110": {"top": 1821942, "left": 8906256, "width": 382768, "height": 146304},
    "Shape 120": {"top": 1821942, "left": 10479024, "width": 1280160, "height": 146304},
    "Shape 121": {"top": 1821942, "left": 10479024, "width": 57607, "height": 146304},
}

for s in slide.shapes:
    if s.name in progress_bar_positions:
        pos = progress_bar_positions[s.name]
        print(f"  {s.name}: top={s.top} -> {pos['top']}, height={s.height} -> {pos['height']}, width={s.width} -> {pos['width']}")
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']

print("\n--- 3. Fixing L section title shapes ---")
l_title_positions = {
    "Shape 131": {"top": 3362706, "left": 4608576, "width": 7461504, "height": 256032},
    "Text 132": {"top": 3399282, "left": 4700016, "width": 7296912, "height": 192024},
}

for s in slide.shapes:
    if s.name in l_title_positions:
        pos = l_title_positions[s.name]
        print(f"  {s.name}: top={s.top} -> {pos['top']}")
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']

print("\n--- 4. Fixing misplaced business line shapes (成事家办, 合伙转介) ---")
biz_fix_positions = {
    "Shape 60": {"top": 3164078, "left": 226568, "width": 36576, "height": 347472},
    "Text 61": {"top": 3164078, "left": 299720, "width": 1243584, "height": 182880},
    "Text 62": {"top": 3346958, "left": 299720, "width": 1243584, "height": 137160},
    "Shape 63": {"top": 3182366, "left": 1579880, "width": 1572768, "height": 274320},
    "Shape 64": {"top": 3182366, "left": 1579880, "width": 1051699, "height": 274320},
    "Text 65": {"top": 3237230, "left": 1616456, "width": 996835, "height": 164592},
    "Shape 66": {"top": 3182366, "left": 2631579, "width": 411119, "height": 274320},
    "Text 67": {"top": 3237230, "left": 2668155, "width": 356255, "height": 164592},
    "Shape 68": {"top": 3182366, "left": 3042698, "width": 109950, "height": 274320},
    "Text 69": {"top": 3228086, "left": 3198368, "width": 749808, "height": 182880},
    
    "Shape 70": {"top": 3808222, "left": 226568, "width": 36576, "height": 347472},
    "Text 71": {"top": 3808222, "left": 299720, "width": 1243584, "height": 182880},
    "Text 72": {"top": 3991102, "left": 299720, "width": 1243584, "height": 137160},
    "Shape 73": {"top": 3826510, "left": 1579880, "width": 1572768, "height": 274320},
    "Shape 74": {"top": 3826510, "left": 1579880, "width": 459202, "height": 274320},
    "Text 75": {"top": 3881374, "left": 1616456, "width": 404338, "height": 164592},
    "Shape 76": {"top": 3826510, "left": 2039082, "width": 1093476, "height": 274320},
    "Text 77": {"top": 3881374, "left": 2075658, "width": 1038612, "height": 164592},
    "Shape 78": {"top": 3826510, "left": 3132558, "width": 20090, "height": 274320},
    "Text 79": {"top": 3872230, "left": 3198368, "width": 749808, "height": 182880},
}

for s in slide.shapes:
    if s.name in biz_fix_positions:
        pos = biz_fix_positions[s.name]
        print(f"  {s.name}: top={s.top} -> {pos['top']}")
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']

print("\n--- 5. Fixing Shape 130 (right-side background) ---")
for s in slide.shapes:
    if s.name == "Shape 130":
        print(f"  Shape 130: top={s.top} -> 3375660, height={s.height} -> 3291840")
        s.top = 3375660
        s.height = 3291840

prs.save(output_path)
print(f"\n✅ Saved fixed PPT to: {output_path}")
