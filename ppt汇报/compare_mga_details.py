from pptx import Presentation
import pandas as pd

def extract_chart_data(chart):
    try:
        categories = []
        if hasattr(chart, 'categories') and chart.categories:
            for cat in chart.categories:
                categories.append(str(cat.label))
        
        series_data = []
        for series in chart.series:
            values = []
            for val in series.values:
                if val is None:
                    values.append('')
                else:
                    try:
                        values.append(str(round(float(val), 2)))
                    except:
                        values.append(str(val))
            series_data.append({
                'name': series.name if series.name else '未命名',
                'values': values
            })
        
        return {
            'categories': categories,
            'series': series_data,
            'has_mga': any('MGA' in str(c).upper() for c in categories) or 
                       any('MGA' in str(s['name']).upper() for s in series_data)
        }
    except Exception as e:
        return {'categories': [], 'series': [], 'has_mga': False, 'error': str(e)}

def extract_slide_details(prs, slide_idx):
    if slide_idx < 1 or slide_idx > len(prs.slides):
        return None
    slide = prs.slides[slide_idx-1]
    
    title = ''
    texts = []
    charts = []
    tables = []
    
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                texts.append(text)
                if not title:
                    title = text[:80]
        
        if shape.has_chart:
            chart_info = extract_chart_data(shape.chart)
            chart_info['name'] = shape.name
            charts.append(chart_info)
        
        if shape.has_table:
            table = shape.table
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip()[:80] for cell in row.cells]
                table_data.append(row_data)
            tables.append(table_data)
    
    return {
        'index': slide_idx,
        'title': title,
        'texts': texts,
        'charts': charts,
        'tables': tables,
        'has_mga': any(c['has_mga'] for c in charts) or any('MGA' in t.upper() for t in texts)
    }

ref_ppt = Presentation('周业绩汇报PPT_W28_20260717.pptx')
test_ppt = Presentation('周业绩汇报PPT_FINAL.pptx')

pages = [1, 4, 5, 7, 8, 9]

print('=' * 120)
print('MGA业务差异分析报告')
print('=' * 120)
print()

for page in pages:
    ref = extract_slide_details(ref_ppt, page)
    test = extract_slide_details(test_ppt, page)
    
    print('-' * 120)
    print('Page {}: {}'.format(page, ref['title'] if ref else '未知'))
    print('-' * 120)
    
    print('【参考文件】周业绩汇报PPT_W28_20260717.pptx')
    print('  包含MGA: {}'.format('是' if ref['has_mga'] else '否'))
    print('  图表数量: {}'.format(len(ref['charts'])))
    for i, ch in enumerate(ref['charts']):
        print('    图表{} [{}]:'.format(i, ch['name']))
        print('      分类: {}'.format(ch['categories']))
        print('      包含MGA: {}'.format('是' if ch['has_mga'] else '否'))
        for s in ch['series']:
            print('      系列[{}]: {}'.format(s['name'], s['values'][:5]))
    
    print()
    print('【测试文件】周业绩汇报PPT_FINAL.pptx')
    print('  包含MGA: {}'.format('是' if test['has_mga'] else '否'))
    print('  图表数量: {}'.format(len(test['charts'])))
    for i, ch in enumerate(test['charts']):
        print('    图表{} [{}]:'.format(i, ch['name']))
        print('      分类: {}'.format(ch['categories']))
        print('      包含MGA: {}'.format('是' if ch['has_mga'] else '否'))
        for s in ch['series']:
            print('      系列[{}]: {}'.format(s['name'], s['values'][:5]))
    
    print()
    if ref['has_mga'] != test['has_mga']:
        if ref['has_mga']:
            print('  ⚠️ 差异: 参考文件包含MGA，但测试文件不包含')
        else:
            print('  ⚠️ 差异: 测试文件包含MGA，但参考文件不包含')
    
    if len(ref['charts']) != len(test['charts']):
        print('  ⚠️ 差异: 图表数量不同 (参考{}个 vs 测试{}个)'.format(len(ref['charts']), len(test['charts'])))
    
    print()
