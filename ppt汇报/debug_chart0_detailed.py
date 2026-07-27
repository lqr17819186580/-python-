import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pptx
from chart_xml_patch import _find_series

def inspect_chart(prs, slide_idx, chart_name, label):
    slide = prs.slides[slide_idx]
    for shape in slide.shapes:
        if shape.name == chart_name and shape.has_chart:
            chart_part = shape.chart.part
            series = _find_series(chart_part._element)
            print(f"\n[{label}] Chart {chart_name} on slide {slide_idx}:")
            for i, ser in enumerate(series):
                tx = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
                name = tx.text if tx is not None else 'Unknown'
                nc = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache')
                pts = nc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt') if nc is not None else []
                print(f"  Series {i}: '{name}', points: {len(pts)}")
            return

print("=== Before update_ppt.py ===")
prs = pptx.Presentation('template.pptx')
inspect_chart(prs, 0, "Chart 0", "Before update")

import update_ppt

print("\n=== After update_ppt.py ===")
prs2 = pptx.Presentation('周业绩汇报PPT_FINAL.pptx')
inspect_chart(prs2, 0, "Chart 0", "After update")