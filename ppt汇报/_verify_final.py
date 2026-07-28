from pptx import Presentation

prs = Presentation('周业绩汇报PPT_FINAL.pptx')
slide7 = prs.slides[6]
print('=== FINAL.pptx 第7页 验证 ===')
print()

# 检查所有长文本框
sowhat_count = 0
for i, shape in enumerate(slide7.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if len(text) > 100:
            sowhat_count += 1
            print(f'[So What] 形状{i}: {shape.name}')
            print(f'  位置: top={shape.top}, left={shape.left}')
            print(f'  文本({len(text)}字): {text[:100]}...')
            print()

print(f'共找到 {sowhat_count} 个 So What 文本框')

# 检查T标题
print()
print('=== T标题检查 ===')
for shape in slide7.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if '签批时效综合分析' in text:
            print(f'T标题: {text}')
            print(f'  位置: top={shape.top}')

# 检查T表格位置
print()
print('=== T表格位置检查 ===')
table_top_values = []
for shape in slide7.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text == '业务线':
            table_top_values.append(shape.top)
            print(f'T表头 "业务线" 位置: top={shape.top}')
            break

# 检查表格数据行位置
for shape in slide7.shapes:
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text == '合计':
            print(f'T表尾 "合计" 位置: top={shape.top}')
            break

print()
if sowhat_count >= 4:
    print('✅ 验证通过: So What 文本框完整')
elif sowhat_count >= 2:
    print('⚠️ 部分通过: 有2个以上So What文本框')
else:
    print('❌ 失败: So What 文本框丢失')