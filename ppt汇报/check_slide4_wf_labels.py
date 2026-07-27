#!/usr/bin/env python3
"""Check Slide 4 waterfall bottom labels"""
from pptx import Presentation

prs = Presentation("周业绩汇报PPT_FINAL_fixed.pptx")
slide4 = prs.slides[3]

print("Slide 4 bottom labels (top around 5660000-5675000):")
for sh in slide4.shapes:
    if sh.has_text_frame and sh.top is not None:
        if 5650000 <= sh.top <= 5680000:
            text = sh.text_frame.text.strip()
            if text:
                try:
                    fill_type = sh.fill.type
                    line_type = sh.line.fill.type
                except:
                    fill_type = "N/A"
                    line_type = "N/A"
                print(f"  name={sh.name!r}, text={text!r}, top={sh.top}, left={sh.left}, fill={fill_type}, line={line_type}")