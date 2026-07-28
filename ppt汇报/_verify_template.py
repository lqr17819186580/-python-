from pptx import Presentation

# 检查 template.pptx 第7页
prs = Presentation('template.pptx')
slide7 = prs.slides[6]
print('=== template.pptx 第7页 ===')
sowhat_count = 0
for i, shape in enumerate(slide7.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if len(text) > 100:
            sowhat_count += 1
            print(f'[So What] 形状{i}: {shape.name}')
            print(f'  位置: top={shape.top}, left={shape.left}')
            print(f'  文本({len(text)}字): {text[:120]}...')
            print()

print(f'template.pptx 第7页共 {sowhat_count} 个 So What 文本框')

# 同时检查 PPT_SOURCE
print()
print('=== PPT_SOURCE 第7页 ===')
prs2 = Presentation('PPT_SOURCE.pptx')
slide7b = prs2.slides[6]
sowhat_count2 = 0
for i, shape in enumerate(slide7b.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if len(text) > 100:
            sowhat_count2 += 1
            print(f'[So What] 形状{i}: {shape.name}')
            print(f'  位置: top={shape.top}, left={shape.left}')
            print(f'  文本({len(text)}字): {text[:120]}...')
            print()

print(f'PPT_SOURCE 第7页共 {sowhat_count2} 个 So What 文本框')