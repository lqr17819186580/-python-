import pptx
from chart_xml_patch import _find_series

prs = pptx.Presentation('周业绩汇报PPT_W28_20260717.pptx')
slide = prs.slides[0]

for shape in slide.shapes:
    if shape.name == 'Chart 0' and shape.has_chart:
        chart_part = shape.chart.part
        series = _find_series(chart_part._element)
        print('参考PPT Chart 0结构:')
        for i, ser in enumerate(series):
            tx = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
            name = tx.text if tx is not None else 'Unknown'
            nc = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache')
            pts = nc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt') if nc is not None else []
            print(f'  Series {i}: "{name}", points: {len(pts)}')
        
        cat = chart_part._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}cat')
        if cat is not None:
            sc = cat.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}strCache')
            if sc is not None:
                ptCount = sc.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}ptCount')
                print(f'  Categories count: {ptCount.get("val") if ptCount is not None else "Unknown"}')
                pts = sc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt')
                for pt in pts:
                    idx = pt.get('idx')
                    v = pt.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
                    print(f'    Category {idx}: "{v.text if v is not None else ""}"')
        break