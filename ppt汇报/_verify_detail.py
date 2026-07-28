from pptx import Presentation

prs = Presentation('template.pptx')
slide7 = prs.slides[6]
print('=== template.pptx 第7页 所有形状 ===')
for i, shape in enumerate(slide7.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text and ('So What' in text or len(text) > 40):
            print(f'形状{i}: {shape.name}')
            print(f'  top={shape.top}, left={shape.left}, w={shape.width}, h={shape.height}')
            print(f'  文本({len(text)}字): {text[:150]}')
            print()

# 检查 "So What" 标签
print('=== So What 标签 ===')
for i, shape in enumerate(slide7.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text.startswith('So What'):
            print(f'形状{i}: {shape.name}')
            print(f'  top={shape.top}, left={shape.left}')
            print(f'  文本: {text}')
            print()