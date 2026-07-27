from pptx import Presentation

target_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL.pptx"

prs = Presentation(target_path)
slide = prs.slides[4]

print("=== Searching for So What text shapes ===")
for s in slide.shapes:
    if s.has_text_frame:
        text = s.text_frame.text.strip()
        if text and len(text) > 20:
            if 'So What' in text or 'So What：' in text:
                print(f"\n  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}")
                print(f"    text='{text[:100]}...'")
        
        if 'BK业务批核' in text or '永明经代目标' in text or '达成率' in text:
            print(f"\n  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}")
            print(f"    text='{text}'")

print("\n=== All shapes with 'BK业务批核' or '永明经代' ===")
for s in slide.shapes:
    if s.has_text_frame:
        text = s.text_frame.text.strip()
        if 'BK业务批核' in text or '永明经代' in text:
            print(f"  {s.name}: top={s.top}, left={s.left}, text='{text[:80]}'")

print("\n=== All shapes in So What area (top 2800000-3500000) ===")
for s in slide.shapes:
    if s.top is not None and 2800000 <= s.top <= 3500000:
        text_content = ""
        if s.has_text_frame:
            text_content = s.text_frame.text.strip()[:80]
        print(f"  {s.name}: top={s.top}, left={s.left}, text='{text_content}'")
