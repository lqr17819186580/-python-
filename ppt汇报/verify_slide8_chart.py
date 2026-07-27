from pptx import Presentation

def verify_chart_data(ppt_path):
    print(f"\n{'='*60}")
    print(f"Verifying Slide 8 charts: {ppt_path}")
    print(f"{'='*60}")
    
    prs = Presentation(ppt_path)
    if len(prs.slides) < 8:
        print(f"ERROR: Only {len(prs.slides)} slides found")
        return
    
    slide = prs.slides[7]  # 0-indexed
    
    for shape in slide.shapes:
        if shape.has_chart:
            chart = shape.chart
            print(f"\nChart: {shape.name}")
            print(f"  Type: {chart.chart_type}")
            
            try:
                # Try to get categories
                categories = []
                for cat in chart.categories:
                    categories.append(cat.label)
                print(f"  Categories: {categories[:5]}... (total {len(categories)})")
            except Exception as e:
                print(f"  Categories error: {e}")
            
            # Get series data
            for i, series in enumerate(chart.series):
                try:
                    values = [v for v in series.values]
                    print(f"  Series [{i}]: {series.name}")
                    print(f"    Values: {values[:5]}... (total {len(values)})")
                except Exception as e:
                    print(f"  Series [{i}] error: {e}")

if __name__ == "__main__":
    verify_chart_data("周业绩汇报PPT_FINAL_MGA.pptx")