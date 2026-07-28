from pptx import Presentation

# 检查template.pptx第7页的So What文本框位置
prs = Presentation('template.pptx')
slide7 = prs.slides[6]
print('=== template.pptx 第7页 So What 文本框详细 ===')

# So What 文本框特征：文本长度 > 100 字
for i, shape in enumerate(slide7.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if len(text) > 100:
            print(f'形状{i}: {shape.name}')
            print(f'  top={shape.top}, left={shape.left}, 宽={shape.width}, 高={shape.height}')
            print(f'  文本前80字: {text[:80]}')
            # 检查是否在T-table删除区域内
            in_delete_zone = (shape.top is not None and shape.left is not None and
                            shape.top >= 3000000 and shape.top < 9500000 and
                            shape.left >= 5000000 and shape.left < 13500000)
            print(f'  在T-table删除区域内: {in_delete_zone}')
            print()

# 检查FINAL.pptx
print('=== FINAL.pptx 第7页 检查 ===')
prs2 = Presentation('周业绩汇报PPT_FINAL.pptx')
slide7_final = prs2.slides[6]
print(f'总形状数: {len(slide7_final.shapes)}')

# 找出所有长文本框
for i, shape in enumerate(slide7_final.shapes):
    if shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if len(text) > 100:
            print(f'形状{i}: {shape.name}')
            print(f'  top={shape.top}, left={shape.left}, 宽={shape.width}, 高={shape.height}')
            print(f'  文本前80字: {text[:80]}')
            print()

# 检查FINAL中是否还有位置在3300000-6500000范围的文本框
print('\n=== FINAL中top在3300000-6500000范围的所有文本框 ===')
for i, shape in enumerate(slide7_final.shapes):
    if shape.has_text_frame and shape.top is not None:
        if 3300000 <= shape.top <= 6500000:
            text = shape.text_frame.text.strip()
            if text:
                print(f'形状{i}: {shape.name}')
                print(f'  top={shape.top}, left={shape.left}')
                print(f'  文本: {text[:60]}...' if len(text) > 60 else f'  文本: {text}')
                print()