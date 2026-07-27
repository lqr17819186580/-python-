#!/usr/bin/env python3
"""Fix three issues:
1. Slide 7 T title missing
2. Slide 4 extra white text box
3. Slide 9 peer data verification
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

IN_PPT = "周业绩汇报PPT_FINAL.pptx"
OUT_PPT = "周业绩汇报PPT_FINAL_fixed.pptx"

def fix_slide7_title(slide):
    print("\n[Slide 7] Checking T title...")
    has_t_title = False
    for sh in slide.shapes:
        if sh.has_text_frame:
            text = sh.text_frame.text.strip()
            if "签批时效" in text or "T" in text:
                has_t_title = True
                print(f"  Found title: {text!r}")
                break
    
    if not has_t_title:
        print("  T title missing, adding...")
        left = Emu(6163056)
        top = Emu(3800000)
        width = Emu(5888736)
        height = Emu(300000)
        
        bg = slide.shapes.add_shape(1, left, top, width, height)
        bg.fill.background()
        bg.line.fill.background()
        
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        run = para.add_run()
        run.text = "T  签批时效综合分析 — 业务线时效 & 分档分布"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        run.font.name = 'Calibri'
        print("  Added T title")

def fix_slide4_extra_textbox(slide):
    print("\n[Slide 4] Checking for extra white text boxes...")
    removed = 0
    for sh in list(slide.shapes):
        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text.strip()
        if not text:
            continue
        try:
            has_white_fill = False
            if sh.fill.type == 1:
                try:
                    if sh.fill.fore_color.rgb == RGBColor(0xFF, 0xFF, 0xFF):
                        has_white_fill = True
                except:
                    pass
            
            has_border = False
            try:
                if sh.line.fill.type != 0:
                    has_border = True
            except:
                pass
            
            if has_white_fill and has_border:
                print(f"  Found suspicious text box: name={sh.name!r}, text={text!r}, top={sh.top}, left={sh.left}")
                sh._element.getparent().remove(sh._element)
                removed += 1
        except Exception as e:
            print(f"  Error checking shape {sh.name}: {e}")
    print(f"  Removed {removed} extra text boxes")

def check_slide9_data(slide):
    print("\n[Slide 9] Checking peer weekly chart data...")
    for sh in slide.shapes:
        if sh.has_chart:
            chart = sh.chart
            print(f"  Chart: {sh.name}")
            for series in chart.series:
                print(f"    Series: {series.name}")
                if len(series.values) >= 28:
                    print(f"      W28 value: {series.values[27]}")
            break

def main():
    print(f"Loading {IN_PPT}...")
    prs = Presentation(IN_PPT)
    
    slide7 = prs.slides[6]
    fix_slide7_title(slide7)
    
    slide4 = prs.slides[3]
    fix_slide4_extra_textbox(slide4)
    
    slide9 = prs.slides[8]
    check_slide9_data(slide9)
    
    prs.save(OUT_PPT)
    print(f"\nSaved to {OUT_PPT}")

if __name__ == "__main__":
    main()