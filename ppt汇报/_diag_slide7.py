from pptx import Presentation
from pptx.util import Emu

prs = Presentation('周业绩汇报PPT_FINAL.pptx')
slide7 = prs.slides[6]
print('=== Slide 7 shapes diagnosis ===')
print(f'Total shapes: {len(slide7.shapes)}')
print()

for i, sh in enumerate(slide7.shapes):
    t = sh.text_frame.text.strip() if sh.has_text_frame else ''
    bg = 'NO_FILL'
    try:
        if sh.fill.type is not None:
            bg = str(sh.fill.type)
            if hasattr(sh.fill, 'fore_color') and sh.fill.fore_color and sh.fill.fore_color.type is not None:
                try:
                    bg += ' rgb=' + str(sh.fill.fore_color.rgb)
                except:
                    pass
    except:
        bg = 'ERR'
    
    left = sh.left if sh.left else 0
    top = sh.top if sh.top else 0
    width = sh.width if sh.width else 0
    height = sh.height if sh.height else 0
    
    is_sowhat = False
    if t.startswith('So What'):
        is_sowhat = True
    elif len(t) > 40 and 6_000_000 <= top <= 7_200_000:
        is_sowhat = True
    
    is_table = hasattr(sh, 'has_table') and sh.has_table
    
    marker = ''
    if is_sowhat:
        marker = ' <<<SOWHAT>>>'
    if is_table:
        marker = ' <<<TABLE>>>'
    
    shape_type = 'table' if is_table else 'shape'
    print(f'[{i:3d}] {shape_type:10s} top={top:>10} left={left:>10} w={width:>10} h={height:>10} bg={bg:30s} text={t[:60]!r}{marker}')