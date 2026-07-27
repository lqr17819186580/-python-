import pptx

prs = pptx.Presentation('周业绩汇报PPT_FINAL.pptx')

print("=== Checking all slides for Chart 0 ===")
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.name == 'Chart 0' and shape.has_chart:
            print(f"Slide {i+1} (slides[{i}]) has Chart 0")