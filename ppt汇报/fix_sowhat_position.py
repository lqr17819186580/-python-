from pptx import Presentation

target_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL.pptx"
output_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL_fixed.pptx"

prs = Presentation(target_path)
slide = prs.slides[4]

print("=== Fixing So What text position ===")

for s in slide.shapes:
    if s.name == "Text 29" and s.left == 4700270:
        print(f"  Found So What paragraph 2: {s.name}")
        print(f"  Current position: top={s.top}, left={s.left}, width={s.width}, height={s.height}")
        print(f"  Target position: top=6467475, left=4700270, width=6459220, height=200025")
        
        s.top = 6467475
        s.left = 4700270
        s.width = 6459220
        s.height = 200025
        
        print(f"  Fixed position: top={s.top}, left={s.left}, width={s.width}, height={s.height}")

prs.save(output_path)
print(f"\n✅ Saved fixed PPT to: {output_path}")
