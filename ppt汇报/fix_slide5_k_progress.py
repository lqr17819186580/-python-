from pptx import Presentation

target_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL.pptx"
output_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL_fixed.pptx"

prs = Presentation(target_path)
slide = prs.slides[4]

print("=== Fixing K section progress bars ===")
progress_bar_positions = {
    "Shape 98": {"top": 1821942, "left": 7333488, "width": 1280160, "height": 146304},
    "Shape 99": {"top": 1821942, "left": 7333488, "width": 839785, "height": 146304},
    "Shape 109": {"top": 1821942, "left": 8906256, "width": 1280160, "height": 146304},
    "Shape 110": {"top": 1821942, "left": 8906256, "width": 382768, "height": 146304},
    "Shape 120": {"top": 1821942, "left": 10479024, "width": 1280160, "height": 146304},
    "Shape 121": {"top": 1821942, "left": 10479024, "width": 57607, "height": 146304},
}

for s in slide.shapes:
    if s.name in progress_bar_positions:
        pos = progress_bar_positions[s.name]
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']
        print(f"  Fixed {s.name}: top={s.top}, height={s.height}")

print("\n=== Fixing K section So-What shapes ===")
so_what_positions = {
    "Shape 126": {"top": 2902077, "left": 4608576, "width": 7461504, "height": 384048},
    "Shape 127": {"top": 2902077, "left": 4608576, "width": 7461504, "height": 27432},
    "Text 128": {"top": 2938653, "left": 4700016, "width": 7296912, "height": 128016},
    "Text 129": {"top": 3057525, "left": 4700016, "width": 7296912, "height": 210312},
}

for s in slide.shapes:
    if s.name in so_what_positions:
        pos = so_what_positions[s.name]
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']
        print(f"  Fixed {s.name}: top={s.top}")

print("\n=== Fixing L section title position ===")
l_title_positions = {
    "Shape 131": {"top": 3362706, "left": 4608576, "width": 7461504, "height": 256032},
    "Text 132": {"top": 3399282, "left": 4700016, "width": 7296912, "height": 192024},
}

for s in slide.shapes:
    if s.name in l_title_positions:
        pos = l_title_positions[s.name]
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']
        print(f"  Fixed {s.name}: top={s.top}")

prs.save(output_path)
print(f"\n✅ Saved fixed PPT to: {output_path}")
