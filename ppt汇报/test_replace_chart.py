import pptx
from chart_updates import slide1_chart_business_type
from data_loader import load_all
from chart_xml_patch import patch_chart

S1, S2, S3, S4 = load_all('S1-总览仪表盘.csv', 'S2-业务端视角.csv', 'S3-执行管理端.csv', 'S4-产品端视角.csv')
cd = slide1_chart_business_type(S1)
cats = [c.label for c in cd.categories]
names = [s.name for s in cd._series]
vals = [[None if v is None else float(v) for v in s.values] for s in cd._series]

print('Input to patch_chart:')
print('Categories:', cats)
print('Series names:', names)
print('Series values:', vals)

prs = pptx.Presentation('周业绩汇报PPT_FINAL.pptx')
slide = prs.slides[0]

for shape in slide.shapes:
    if shape.name == 'Chart 0':
        chart_part = shape.chart.part
        
        print('\nBefore patch:')
        for i, ser in enumerate(chart_part._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}ser')):
            tx = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
            name = tx.text if tx is not None else 'Unknown'
            nc = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache')
            pts = nc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt') if nc is not None else []
            print('Series', i, ':', name, ', points:', len(pts))
        
        patch_chart(chart_part, vals, categories=cats, series_names=names)
        
        print('\nAfter patch:')
        for i, ser in enumerate(chart_part._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}ser')):
            tx = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
            name = tx.text if tx is not None else 'Unknown'
            nc = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache')
            pts = nc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt') if nc is not None else []
            print('Series', i, ':', name, ', points:', len(pts))
        
        prs.save('test_replace_output.pptx')
        print('\nSaved test_replace_output.pptx')
        
        print('\n=== Checking output file ===')
        prs2 = pptx.Presentation('test_replace_output.pptx')
        slide2 = prs2.slides[0]
        for shape in slide2.shapes:
            if shape.name == 'Chart 0':
                chart_part2 = shape.chart.part
                print('After save and reopen:')
                for i, ser in enumerate(chart_part2._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}ser')):
                    tx = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}v')
                    name = tx.text if tx is not None else 'Unknown'
                    nc = ser.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}numCache')
                    pts = nc.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}pt') if nc is not None else []
                    print('Series', i, ':', name, ', points:', len(pts))
