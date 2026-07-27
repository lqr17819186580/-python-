#!/usr/bin/env python3
"""Check original PPT for MGA batch label"""
from pptx import Presentation

prs = Presentation("周业绩汇报PPT_FINAL.pptx")
slide4 = prs.slides[3]

print("All shapes containing 'MGA':")
for sh in slide4.shapes:
    if sh.has_text_frame:
        text = sh.text_frame.text.strip()
        if "MGA" in text:
            try:
                fill_type = sh.fill.type
                line_type = sh.line.fill.type
            except:
                fill_type = "N/A"
                line_type = "N/A"
            print(f"  name={sh.name!r}, text={text!r}, top={sh.top}, left={sh.left}, fill={fill_type}, line={line_type}")