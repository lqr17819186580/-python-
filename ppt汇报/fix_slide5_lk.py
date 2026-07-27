from pptx import Presentation
from pptx.util import Emu

template_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\template.pptx"
target_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL.pptx"
output_path = r"D:\数据中台支持\业绩整合自动化脚本\ppt汇报\周业绩汇报PPT_FINAL_fixed.pptx"

print("=== Analyzing template.pptx K section progress bars ===")
template_prs = Presentation(template_path)
template_slide = template_prs.slides[4]

template_progress_shapes = {}
for s in template_slide.shapes:
    if s.top is not None and 850000 <= s.top <= 2800000:
        if s.name.startswith('Shape') and not s.has_text_frame:
            fill_color = None
            if hasattr(s, 'fill'):
                try:
                    fill = s.fill
                    if fill.type == 1:
                        fill_color = fill.fore_color.rgb
                except:
                    pass
            template_progress_shapes[s.name] = {
                'top': s.top,
                'left': s.left,
                'width': s.width,
                'height': s.height,
                'fill': fill_color
            }
            print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, fill={fill_color}")

print("\n=== Analyzing target PPT K section progress bars ===")
prs = Presentation(target_path)
slide = prs.slides[4]

target_progress_shapes = {}
for s in slide.shapes:
    if s.top is not None and 850000 <= s.top <= 2800000:
        if s.name.startswith('Shape') and not s.has_text_frame:
            fill_color = None
            if hasattr(s, 'fill'):
                try:
                    fill = s.fill
                    if fill.type == 1:
                        fill_color = fill.fore_color.rgb
                except:
                    pass
            target_progress_shapes[s.name] = {
                'top': s.top,
                'left': s.left,
                'width': s.width,
                'height': s.height,
                'fill': fill_color
            }
            print(f"  {s.name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}, fill={fill_color}")

print("\n=== Fixing K section progress bars ===")
for shape_name, template_info in template_progress_shapes.items():
    for s in slide.shapes:
        if s.name == shape_name:
            s.top = template_info['top']
            s.left = template_info['left']
            s.width = template_info['width']
            s.height = template_info['height']
            print(f"  Fixed {shape_name}: top={s.top}, left={s.left}, width={s.width}, height={s.height}")
            break

print("\n=== Fixing L section title position ===")
l_title_positions = {
    "Shape 131": {"top": 3362706, "left": 4608576, "width": 7461504, "height": 256032},
    "Text 132": {"top": 3399282, "left": 4700016, "width": 7296912, "height": 192024},
}

for s in slide.shapes:
    if s.name in l_title_positions:
        pos = l_title_positions[s.name]
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']
        print(f"  Fixed {s.name}: top={s.top}")

print("\n=== Fixing K section So-What shapes ===")
so_what_positions = {
    "Shape 126": {"top": 2902077, "left": 4608576, "width": 7461504, "height": 384048},
    "Shape 127": {"top": 2902077, "left": 4608576, "width": 7461504, "height": 27432},
    "Text 128": {"top": 2938653, "left": 4700016, "width": 7296912, "height": 128016},
    "Text 129": {"top": 3057525, "left": 4700016, "width": 7296912, "height": 210312},
}

for s in slide.shapes:
    if s.name in so_what_positions:
        pos = so_what_positions[s.name]
        s.top = pos['top']
        s.left = pos['left']
        s.width = pos['width']
        s.height = pos['height']
        print(f"  Fixed {s.name}: top={s.top}")

prs.save(output_path)
print(f"\n✅ Saved fixed PPT to: {output_path}")
